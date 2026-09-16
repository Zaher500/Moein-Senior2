# from pathlib import Path
# import shutil
# import subprocess

# import fitz  # PyMuPDF
# from docx import Document
# from pptx import Presentation

# from .schemas import ProcessedDocument


# RASTER_IMAGE_EXTENSIONS = {
#     ".png",
#     ".jpg",
#     ".jpeg",
#     ".bmp",
#     ".tif",
#     ".tiff",
#     ".webp",
# }


# def extract_candidate_images(
#     document: ProcessedDocument,
#     output_dir: str,
#     zoom: float = 2.0,
# ) -> list[str]:
#     """
#     Extract images marked with:

#         process_with_ai = True

#     Supported document types:
#     - PDF
#     - DOCX
#     - PPTX

#     PDF:
#         Images are rendered from their page bounding boxes
#         exactly as they appear on the PDF page.

#     DOCX:
#         Original embedded image bytes are extracted directly
#         from the Word document.

#     PPTX:
#         Raster pictures are extracted directly.

#         Vector pictures such as WMF / EMF, charts,
#         and grouped diagrams are rendered from the slide
#         and cropped according to their PowerPoint bounding box.

#     Returns:
#         List of generated image file paths.
#     """

#     # =========================================================
#     # Route according to document type
#     # =========================================================

#     if document.file_type == "pdf":

#         return _extract_pdf_candidate_images(
#             document=document,
#             output_dir=output_dir,
#             zoom=zoom,
#         )

#     if document.file_type == "docx":

#         return _extract_docx_candidate_images(
#             document=document,
#             output_dir=output_dir,
#         )

#     if document.file_type == "pptx":

#         return _extract_pptx_candidate_images(
#             document=document,
#             output_dir=output_dir,
#             zoom=zoom,
#         )

#     raise ValueError(
#         f"extract_candidate_images does not support "
#         f"file type: {document.file_type}"
#     )


# # =============================================================
# # PDF
# # =============================================================


# def _extract_pdf_candidate_images(
#     document: ProcessedDocument,
#     output_dir: str,
#     zoom: float = 2.0,
# ) -> list[str]:
#     """
#     Extract PDF images marked with:

#         process_with_ai = True

#     Images are rendered from their page bounding boxes,
#     exactly as they appear on the PDF page.

#     This preserves the existing PDF behavior.
#     """

#     output_path = Path(
#         output_dir
#     )

#     output_path.mkdir(
#         parents=True,
#         exist_ok=True,
#     )

#     extracted_files = []

#     try:

#         with fitz.open(
#             document.file_path
#         ) as pdf:

#             for unit in (
#                 document.get_ordered_units()
#             ):

#                 # PDF pages in PyMuPDF are zero-based,
#                 # while our unit.index is one-based.
#                 page_index = (
#                     unit.index - 1
#                 )

#                 if (
#                     page_index < 0
#                     or page_index >= pdf.page_count
#                 ):
#                     continue

#                 page = pdf[
#                     page_index
#                 ]

#                 for element in (
#                     unit.get_ordered_elements()
#                 ):

#                     if element.type != "image":
#                         continue

#                     if not element.metadata.get(
#                         "process_with_ai",
#                         False,
#                     ):
#                         continue

#                     if element.bbox is None:
#                         continue

#                     bbox = element.bbox

#                     rect = fitz.Rect(
#                         bbox.x,
#                         bbox.y,
#                         bbox.x + bbox.width,
#                         bbox.y + bbox.height,
#                     )

#                     # Make sure the crop stays
#                     # inside the PDF page.
#                     rect = (
#                         rect
#                         & page.rect
#                     )

#                     if rect.is_empty:
#                         continue

#                     matrix = fitz.Matrix(
#                         zoom,
#                         zoom,
#                     )

#                     pixmap = (
#                         page.get_pixmap(
#                             matrix=matrix,
#                             clip=rect,
#                             alpha=False,
#                         )
#                     )

#                     image_filename = (
#                         f"page_{unit.index:03d}"
#                         f"_element_"
#                         f"{element.order:03d}.png"
#                     )

#                     image_path = (
#                         output_path
#                         / image_filename
#                     )

#                     pixmap.save(
#                         str(image_path)
#                     )

#                     element.metadata[
#                         "extracted_image_path"
#                     ] = str(
#                         image_path
#                     )

#                     element.metadata[
#                         "render_zoom"
#                     ] = zoom

#                     extracted_files.append(
#                         str(image_path)
#                     )

#     except Exception as e:

#         raise RuntimeError(
#             f"Failed to extract PDF images from "
#             f"'{document.file_path}': {str(e)}"
#         ) from e

#     return extracted_files


# # =============================================================
# # DOCX
# # =============================================================


# def _extract_docx_candidate_images(
#     document: ProcessedDocument,
#     output_dir: str,
# ) -> list[str]:
#     """
#     Extract DOCX embedded images marked with:

#         process_with_ai = True

#     DOCX images are not rendered from page coordinates.

#     Instead, the original image bytes stored inside
#     the Word document are extracted directly.
#     """

#     output_path = Path(
#         output_dir
#     )

#     output_path.mkdir(
#         parents=True,
#         exist_ok=True,
#     )

#     extracted_files = []

#     try:

#         doc = Document(
#             document.file_path
#         )

#         for unit in (
#             document.get_ordered_units()
#         ):

#             for element in (
#                 unit.get_ordered_elements()
#             ):

#                 if element.type != "image":
#                     continue

#                 if not element.metadata.get(
#                     "process_with_ai",
#                     False,
#                 ):
#                     continue

#                 relationship_id = (
#                     element.metadata.get(
#                         "relationship_id"
#                     )
#                 )

#                 image_partname = (
#                     element.metadata.get(
#                         "image_partname"
#                     )
#                 )

#                 image_part = None

#                 # =================================================
#                 # 1. Find image using relationship_id
#                 # =================================================

#                 if relationship_id:

#                     try:

#                         image_part = (
#                             doc.part.related_parts[
#                                 relationship_id
#                             ]
#                         )

#                     except Exception:

#                         image_part = None

#                 # =================================================
#                 # 2. Fallback: find image using part name
#                 # =================================================

#                 if (
#                     image_part is None
#                     and image_partname
#                 ):

#                     for related_part in (
#                         doc.part.related_parts.values()
#                     ):

#                         current_partname = str(
#                             getattr(
#                                 related_part,
#                                 "partname",
#                                 "",
#                             )
#                         )

#                         if (
#                             current_partname
#                             == image_partname
#                         ):

#                             image_part = (
#                                 related_part
#                             )

#                             break

#                 # =================================================
#                 # 3. Image could not be located
#                 # =================================================

#                 if image_part is None:

#                     element.metadata[
#                         "extraction_status"
#                     ] = (
#                         "failed_image_not_found"
#                     )

#                     continue

#                 # =================================================
#                 # 4. Read original image bytes
#                 # =================================================

#                 image_blob = getattr(
#                     image_part,
#                     "blob",
#                     None,
#                 )

#                 if not image_blob:

#                     element.metadata[
#                         "extraction_status"
#                     ] = (
#                         "failed_empty_image_blob"
#                     )

#                     continue

#                 # =================================================
#                 # 5. Determine original image extension
#                 # =================================================

#                 partname = str(
#                     getattr(
#                         image_part,
#                         "partname",
#                         "",
#                     )
#                 )

#                 extension = (
#                     Path(
#                         partname
#                     ).suffix.lower()
#                 )

#                 if not extension:

#                     extension = ".png"

#                 # =================================================
#                 # 6. Save original image
#                 # =================================================

#                 image_filename = (
#                     f"unit_{unit.index:03d}"
#                     f"_element_"
#                     f"{element.order:03d}"
#                     f"{extension}"
#                 )

#                 image_path = (
#                     output_path
#                     / image_filename
#                 )

#                 with open(
#                     image_path,
#                     "wb",
#                 ) as image_file:

#                     image_file.write(
#                         image_blob
#                     )

#                 # =================================================
#                 # 7. Save extraction result into IR
#                 # =================================================

#                 element.metadata[
#                     "extracted_image_path"
#                 ] = str(
#                     image_path
#                 )

#                 element.metadata[
#                     "extraction_status"
#                 ] = "done"

#                 element.metadata[
#                     "extracted_extension"
#                 ] = extension

#                 extracted_files.append(
#                     str(image_path)
#                 )

#     except Exception as e:

#         raise RuntimeError(
#             f"Failed to extract DOCX images from "
#             f"'{document.file_path}': {str(e)}"
#         ) from e

#     return extracted_files


# # =============================================================
# # PPTX
# # =============================================================


# def _extract_pptx_candidate_images(
#     document: ProcessedDocument,
#     output_dir: str,
#     zoom: float = 2.0,
# ) -> list[str]:
#     """
#     Extract PPTX visual elements marked with:

#         process_with_ai = True

#     Strategy:

#     1. Normal raster pictures:
#        Extract their original embedded bytes directly.

#     2. Vector pictures such as WMF / EMF:
#        Render the PowerPoint slide and crop the picture area.

#     3. Charts and grouped diagrams:
#        Render the slide and crop the corresponding shape area.

#     Slide rendering is performed through LibreOffice.
#     """

#     output_path = Path(
#         output_dir
#     )

#     output_path.mkdir(
#         parents=True,
#         exist_ok=True,
#     )

#     extracted_files = []

#     try:

#         presentation = Presentation(
#             document.file_path
#         )

#         candidates = []

#         # =====================================================
#         # 1. Collect visual candidates
#         # =====================================================

#         for unit in (
#             document.get_ordered_units()
#         ):

#             for element in (
#                 unit.get_ordered_elements()
#             ):

#                 if element.type != "image":
#                     continue

#                 if not element.metadata.get(
#                     "process_with_ai",
#                     False,
#                 ):
#                     continue

#                 candidates.append(
#                     (
#                         unit,
#                         element,
#                     )
#                 )

#         # =====================================================
#         # 2. Determine whether slide rendering is required
#         # =====================================================

#         needs_slide_rendering = False

#         for unit, element in candidates:

#             visual_type = (
#                 element.metadata.get(
#                     "visual_type",
#                     ""
#                 )
#             )

#             render_mode = (
#                 element.metadata.get(
#                     "render_mode",
#                     ""
#                 )
#             )

#             extension = (
#                 _normalize_extension(
#                     element.metadata.get(
#                         "image_extension"
#                     )
#                 )
#             )

#             if (
#                 render_mode == "slide_crop"
#                 or visual_type in {
#                     "chart",
#                     "group",
#                 }
#                 or (
#                     visual_type == "picture"
#                     and extension
#                     not in RASTER_IMAGE_EXTENSIONS
#                 )
#             ):

#                 needs_slide_rendering = True

#                 break

#         # =====================================================
#         # 3. Render PPTX to PDF if required
#         # =====================================================

#         rendered_pdf_path = None
#         render_error = None

#         if needs_slide_rendering:

#             try:

#                 rendered_pdf_path = (
#                     _render_pptx_to_pdf(
#                         pptx_path=document.file_path,
#                         output_dir=output_path,
#                     )
#                 )

#             except Exception as e:

#                 render_error = str(
#                     e
#                 )

#         # =====================================================
#         # 4. Open rendered presentation when available
#         # =====================================================

#         rendered_pdf = None

#         if rendered_pdf_path:

#             rendered_pdf = fitz.open(
#                 str(
#                     rendered_pdf_path
#                 )
#             )

#         try:

#             # =================================================
#             # 5. Process candidates
#             # =================================================

#             for unit, element in candidates:

#                 slide_index = (
#                     unit.index - 1
#                 )

#                 if (
#                     slide_index < 0
#                     or slide_index
#                     >= len(
#                         presentation.slides
#                     )
#                 ):

#                     element.metadata[
#                         "extraction_status"
#                     ] = (
#                         "failed_slide_not_found"
#                     )

#                     continue

#                 slide = (
#                     presentation.slides[
#                         slide_index
#                     ]
#                 )

#                 shape_id = (
#                     element.metadata.get(
#                         "shape_id"
#                     )
#                 )

#                 shape = _find_pptx_shape_by_id(
#                     slide=slide,
#                     shape_id=shape_id,
#                 )

#                 if shape is None:

#                     element.metadata[
#                         "extraction_status"
#                     ] = (
#                         "failed_shape_not_found"
#                     )

#                     continue

#                 visual_type = (
#                     element.metadata.get(
#                         "visual_type",
#                         ""
#                     )
#                 )

#                 render_mode = (
#                     element.metadata.get(
#                         "render_mode",
#                         ""
#                     )
#                 )

#                 # =============================================
#                 # 5A. Embedded picture
#                 # =============================================

#                 if (
#                     visual_type == "picture"
#                     and render_mode == "embedded"
#                 ):

#                     extension = (
#                         _get_pptx_picture_extension(
#                             shape=shape,
#                             element=element,
#                         )
#                     )

#                     # -----------------------------------------
#                     # Raster image
#                     # -----------------------------------------

#                     if (
#                         extension
#                         in RASTER_IMAGE_EXTENSIONS
#                     ):

#                         image_blob = (
#                             _get_pptx_picture_blob(
#                                 shape
#                             )
#                         )

#                         if not image_blob:

#                             element.metadata[
#                                 "extraction_status"
#                             ] = (
#                                 "failed_empty_image_blob"
#                             )

#                             continue

#                         image_filename = (
#                             f"slide_"
#                             f"{unit.index:03d}"
#                             f"_element_"
#                             f"{element.order:03d}"
#                             f"{extension}"
#                         )

#                         image_path = (
#                             output_path
#                             / image_filename
#                         )

#                         with open(
#                             image_path,
#                             "wb",
#                         ) as image_file:

#                             image_file.write(
#                                 image_blob
#                             )

#                         element.metadata[
#                             "extracted_image_path"
#                         ] = str(
#                             image_path
#                         )

#                         element.metadata[
#                             "extraction_status"
#                         ] = "done"

#                         element.metadata[
#                             "extracted_extension"
#                         ] = extension

#                         element.metadata[
#                             "extraction_strategy"
#                         ] = (
#                             "embedded_original"
#                         )

#                         extracted_files.append(
#                             str(image_path)
#                         )

#                         continue

#                     # -----------------------------------------
#                     # Vector picture:
#                     #
#                     # WMF / EMF / SVG / etc.
#                     #
#                     # Render slide and crop it.
#                     # -----------------------------------------

#                     success = (
#                         _extract_pptx_slide_crop(
#                             document=document,
#                             unit=unit,
#                             element=element,
#                             rendered_pdf=rendered_pdf,
#                             output_path=output_path,
#                             zoom=zoom,
#                             original_extension=extension,
#                             render_error=render_error,
#                         )
#                     )

#                     if success:

#                         extracted_files.append(
#                             element.metadata[
#                                 "extracted_image_path"
#                             ]
#                         )

#                     continue

#                 # =============================================
#                 # 5B. Chart / grouped diagram / slide crop
#                 # =============================================

#                 if (
#                     render_mode == "slide_crop"
#                     or visual_type in {
#                         "chart",
#                         "group",
#                     }
#                 ):

#                     success = (
#                         _extract_pptx_slide_crop(
#                             document=document,
#                             unit=unit,
#                             element=element,
#                             rendered_pdf=rendered_pdf,
#                             output_path=output_path,
#                             zoom=zoom,
#                             original_extension=None,
#                             render_error=render_error,
#                         )
#                     )

#                     if success:

#                         extracted_files.append(
#                             element.metadata[
#                                 "extracted_image_path"
#                             ]
#                         )

#                     continue

#                 # =============================================
#                 # Unknown PPTX visual type
#                 # =============================================

#                 element.metadata[
#                     "extraction_status"
#                 ] = (
#                     "failed_unsupported_visual_type"
#                 )

#         finally:

#             if rendered_pdf is not None:

#                 rendered_pdf.close()

#     except Exception as e:

#         raise RuntimeError(
#             f"Failed to extract PPTX images from "
#             f"'{document.file_path}': {str(e)}"
#         ) from e

#     return extracted_files


# def _extract_pptx_slide_crop(
#     document: ProcessedDocument,
#     unit,
#     element,
#     rendered_pdf,
#     output_path: Path,
#     zoom: float,
#     original_extension=None,
#     render_error=None,
# ) -> bool:
#     """
#     Crop one PPTX visual element from the rendered slide.

#     PPTX coordinates are stored in EMU.

#     The rendered PDF page uses PDF points.

#     Because slide dimensions and element bounding boxes use
#     the same PPTX coordinate system, conversion is performed
#     using ratios rather than fixed unit conversion.
#     """

#     if rendered_pdf is None:

#         element.metadata[
#             "extraction_status"
#         ] = (
#             "failed_slide_rendering_unavailable"
#         )

#         if render_error:

#             element.metadata[
#                 "extraction_error"
#             ] = render_error

#         return False

#     slide_index = (
#         unit.index - 1
#     )

#     if (
#         slide_index < 0
#         or slide_index
#         >= rendered_pdf.page_count
#     ):

#         element.metadata[
#             "extraction_status"
#         ] = (
#             "failed_rendered_slide_not_found"
#         )

#         return False

#     if element.bbox is None:

#         element.metadata[
#             "extraction_status"
#         ] = (
#             "failed_missing_bbox"
#         )

#         return False

#     slide_width = float(
#         unit.metadata.get(
#             "width",
#             0,
#         )
#     )

#     slide_height = float(
#         unit.metadata.get(
#             "height",
#             0,
#         )
#     )

#     if (
#         slide_width <= 0
#         or slide_height <= 0
#     ):

#         element.metadata[
#             "extraction_status"
#         ] = (
#             "failed_invalid_slide_dimensions"
#         )

#         return False

#     page = rendered_pdf[
#         slide_index
#     ]

#     bbox = element.bbox

#     # =========================================================
#     # Convert PPTX bbox to ratios
#     # =========================================================

#     x_ratio = (
#         bbox.x
#         / slide_width
#     )

#     y_ratio = (
#         bbox.y
#         / slide_height
#     )

#     width_ratio = (
#         bbox.width
#         / slide_width
#     )

#     height_ratio = (
#         bbox.height
#         / slide_height
#     )

#     # =========================================================
#     # Convert ratios into rendered PDF coordinates
#     # =========================================================

#     x0 = (
#         page.rect.x0
#         + (
#             x_ratio
#             * page.rect.width
#         )
#     )

#     y0 = (
#         page.rect.y0
#         + (
#             y_ratio
#             * page.rect.height
#         )
#     )

#     x1 = (
#         x0
#         + (
#             width_ratio
#             * page.rect.width
#         )
#     )

#     y1 = (
#         y0
#         + (
#             height_ratio
#             * page.rect.height
#         )
#     )

#     rect = fitz.Rect(
#         x0,
#         y0,
#         x1,
#         y1,
#     )

#     rect = (
#         rect
#         & page.rect
#     )

#     if rect.is_empty:

#         element.metadata[
#             "extraction_status"
#         ] = (
#             "failed_empty_render_crop"
#         )

#         return False

#     matrix = fitz.Matrix(
#         zoom,
#         zoom,
#     )

#     pixmap = page.get_pixmap(
#         matrix=matrix,
#         clip=rect,
#         alpha=False,
#     )

#     image_filename = (
#         f"slide_{unit.index:03d}"
#         f"_element_"
#         f"{element.order:03d}.png"
#     )

#     image_path = (
#         output_path
#         / image_filename
#     )

#     pixmap.save(
#         str(image_path)
#     )

#     # =========================================================
#     # Store extraction information
#     # =========================================================

#     element.metadata[
#         "extracted_image_path"
#     ] = str(
#         image_path
#     )

#     element.metadata[
#         "extraction_status"
#     ] = "done"

#     element.metadata[
#         "extracted_extension"
#     ] = ".png"

#     element.metadata[
#         "extraction_strategy"
#     ] = "slide_crop"

#     element.metadata[
#         "render_zoom"
#     ] = zoom

#     if original_extension:

#         element.metadata[
#             "original_extension"
#         ] = original_extension

#     element.metadata.pop(
#         "extraction_error",
#         None,
#     )

#     return True


# def _find_pptx_shape_by_id(
#     slide,
#     shape_id,
# ):
#     """
#     Find a top-level PowerPoint shape using its shape_id.
#     """

#     if shape_id is None:
#         return None

#     for shape in slide.shapes:

#         if (
#             shape.shape_id
#             == shape_id
#         ):

#             return shape

#     return None


# def _get_pptx_picture_blob(
#     shape,
# ):
#     """
#     Safely return the embedded picture bytes from
#     a normal Picture or PlaceholderPicture.
#     """

#     try:

#         image = shape.image

#         return image.blob

#     except Exception:

#         return None


# def _get_pptx_picture_extension(
#     shape,
#     element,
# ) -> str:
#     """
#     Determine a normalized PowerPoint picture extension.

#     Example:

#         png  -> .png
#         jpg  -> .jpg
#         wmf  -> .wmf
#     """

#     extension = ""

#     try:

#         image = shape.image

#         extension = getattr(
#             image,
#             "ext",
#             "",
#         )

#     except Exception:

#         extension = ""

#     if not extension:

#         extension = (
#             element.metadata.get(
#                 "image_extension",
#                 "",
#             )
#         )

#     return _normalize_extension(
#         extension
#     )


# def _normalize_extension(
#     extension,
# ) -> str:
#     """
#     Normalize image extensions into lowercase
#     strings beginning with a dot.
#     """

#     if not extension:
#         return ""

#     extension = str(
#         extension
#     ).strip().lower()

#     if not extension:
#         return ""

#     if not extension.startswith(
#         "."
#     ):

#         extension = (
#             "."
#             + extension
#         )

#     return extension


# def _render_pptx_to_pdf(
#     pptx_path: str,
#     output_dir: Path,
# ) -> Path:
#     """
#     Render the PPTX presentation to PDF using LibreOffice.

#     This allows:
#     - WMF / EMF pictures
#     - charts
#     - grouped diagrams

#     to be converted into normal raster crops later.

#     The function searches:
#     1. PATH
#     2. Common Windows LibreOffice locations
#     """

#     libreoffice = (
#         _find_libreoffice_executable()
#     )

#     if libreoffice is None:

#         raise RuntimeError(
#             "LibreOffice was not found. "
#             "PPTX vector pictures, charts, and grouped "
#             "diagrams require slide rendering before "
#             "OCR / Vision."
#         )

#     render_dir = (
#         output_dir
#         / "_pptx_render"
#     )

#     render_dir.mkdir(
#         parents=True,
#         exist_ok=True,
#     )

#     command = [
#         str(libreoffice),
#         "--headless",
#         "--convert-to",
#         "pdf",
#         "--outdir",
#         str(render_dir),
#         str(
#             Path(
#                 pptx_path
#             ).resolve()
#         ),
#     ]

#     result = subprocess.run(
#         command,
#         capture_output=True,
#         text=True,
#         timeout=120,
#     )

#     expected_pdf = (
#         render_dir
#         / (
#             Path(
#                 pptx_path
#             ).stem
#             + ".pdf"
#         )
#     )

#     if (
#         result.returncode != 0
#         or not expected_pdf.exists()
#     ):

#         stdout = (
#             result.stdout
#             or ""
#         ).strip()

#         stderr = (
#             result.stderr
#             or ""
#         ).strip()

#         raise RuntimeError(
#             "LibreOffice failed to render PPTX. "
#             f"stdout='{stdout}' "
#             f"stderr='{stderr}'"
#         )

#     return expected_pdf


# def _find_libreoffice_executable():
#     """
#     Find LibreOffice / soffice executable.

#     Supports:
#     - Linux/macOS PATH
#     - Windows PATH
#     - common Windows installation locations
#     """

#     # =========================================================
#     # PATH
#     # =========================================================

#     for command in (
#         "soffice",
#         "libreoffice",
#     ):

#         executable = shutil.which(
#             command
#         )

#         if executable:

#             return Path(
#                 executable
#             )

#     # =========================================================
#     # Common Windows locations
#     # =========================================================

#     windows_candidates = [
#         Path(
#             r"C:\Program Files\LibreOffice"
#             r"\program\soffice.exe"
#         ),
#         Path(
#             r"C:\Program Files (x86)\LibreOffice"
#             r"\program\soffice.exe"
#         ),
#     ]

#     for candidate in windows_candidates:

#         if candidate.exists():

#             return candidate

#     return None









from pathlib import Path
import hashlib
import shutil
import subprocess

import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation

from .schemas import ProcessedDocument


RASTER_IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".bmp",
    ".tif",
    ".tiff",
    ".webp",
}


def extract_candidate_images(
    document: ProcessedDocument,
    output_dir: str,
    zoom: float = 2.0,
) -> list[str]:
    """
    Extract images marked with:

        process_with_ai = True

    Supported document types:
    - PDF
    - DOCX
    - PPTX

    PDF:
        Images are rendered from their page bounding boxes
        exactly as they appear on the PDF page.

    DOCX:
        Original embedded image bytes are extracted directly
        from the Word document.

    PPTX:
        Raster pictures are extracted directly.

        Vector pictures such as WMF / EMF, charts,
        and grouped diagrams are rendered from the slide
        and cropped according to their PowerPoint bounding box.

    Returns:
        List of generated image file paths.
    """

    # =========================================================
    # Route according to document type
    # =========================================================

    if document.file_type == "pdf":

        return _extract_pdf_candidate_images(
            document=document,
            output_dir=output_dir,
            zoom=zoom,
        )

    if document.file_type == "docx":

        return _extract_docx_candidate_images(
            document=document,
            output_dir=output_dir,
        )

    if document.file_type == "pptx":

        return _extract_pptx_candidate_images(
            document=document,
            output_dir=output_dir,
            zoom=zoom,
        )

    raise ValueError(
        f"extract_candidate_images does not support "
        f"file type: {document.file_type}"
    )


# =============================================================
# PDF
# =============================================================


def _extract_pdf_candidate_images(
    document: ProcessedDocument,
    output_dir: str,
    zoom: float = 2.0,
) -> list[str]:
    """
    Extract PDF images marked with:

        process_with_ai = True

    Images are rendered from their page bounding boxes,
    exactly as they appear on the PDF page.

    This preserves the existing PDF behavior.
    """

    output_path = Path(
        output_dir
    )

    output_path.mkdir(
        parents=True,
        exist_ok=True,
    )

    extracted_files = []

    try:

        with fitz.open(
            document.file_path
        ) as pdf:

            for unit in (
                document.get_ordered_units()
            ):

                # PDF pages in PyMuPDF are zero-based,
                # while our unit.index is one-based.
                page_index = (
                    unit.index - 1
                )

                if (
                    page_index < 0
                    or page_index >= pdf.page_count
                ):
                    continue

                page = pdf[
                    page_index
                ]

                for element in (
                    unit.get_ordered_elements()
                ):

                    if element.type != "image":
                        continue

                    if not element.metadata.get(
                        "process_with_ai",
                        False,
                    ):
                        continue

                    if element.bbox is None:
                        continue

                    bbox = element.bbox

                    rect = fitz.Rect(
                        bbox.x,
                        bbox.y,
                        bbox.x + bbox.width,
                        bbox.y + bbox.height,
                    )

                    # Make sure the crop stays
                    # inside the PDF page.
                    rect = (
                        rect
                        & page.rect
                    )

                    if rect.is_empty:
                        continue

                    matrix = fitz.Matrix(
                        zoom,
                        zoom,
                    )

                    pixmap = (
                        page.get_pixmap(
                            matrix=matrix,
                            clip=rect,
                            alpha=False,
                        )
                    )

                    image_filename = (
                        f"page_{unit.index:03d}"
                        f"_element_"
                        f"{element.order:03d}.png"
                    )

                    image_path = (
                        output_path
                        / image_filename
                    )

                    pixmap.save(
                        str(image_path)
                    )

                    element.metadata[
                        "extracted_image_path"
                    ] = str(
                        image_path
                    )

                    element.metadata[
                        "extracted_image_hash"
                    ] = _calculate_extracted_image_hash(
                        image_path
                    )

                    element.metadata[
                        "render_zoom"
                    ] = zoom

                    extracted_files.append(
                        str(image_path)
                    )

    except Exception as e:

        raise RuntimeError(
            f"Failed to extract PDF images from "
            f"'{document.file_path}': {str(e)}"
        ) from e

    return extracted_files


# =============================================================
# DOCX
# =============================================================


def _extract_docx_candidate_images(
    document: ProcessedDocument,
    output_dir: str,
) -> list[str]:
    """
    Extract DOCX embedded images marked with:

        process_with_ai = True

    DOCX images are not rendered from page coordinates.

    Instead, the original image bytes stored inside
    the Word document are extracted directly.
    """

    output_path = Path(
        output_dir
    )

    output_path.mkdir(
        parents=True,
        exist_ok=True,
    )

    extracted_files = []

    try:

        doc = Document(
            document.file_path
        )

        for unit in (
            document.get_ordered_units()
        ):

            for element in (
                unit.get_ordered_elements()
            ):

                if element.type != "image":
                    continue

                if not element.metadata.get(
                    "process_with_ai",
                    False,
                ):
                    continue

                relationship_id = (
                    element.metadata.get(
                        "relationship_id"
                    )
                )

                image_partname = (
                    element.metadata.get(
                        "image_partname"
                    )
                )

                image_part = None

                # =================================================
                # 1. Find image using relationship_id
                # =================================================

                if relationship_id:

                    try:

                        image_part = (
                            doc.part.related_parts[
                                relationship_id
                            ]
                        )

                    except Exception:

                        image_part = None

                # =================================================
                # 2. Fallback: find image using part name
                # =================================================

                if (
                    image_part is None
                    and image_partname
                ):

                    for related_part in (
                        doc.part.related_parts.values()
                    ):

                        current_partname = str(
                            getattr(
                                related_part,
                                "partname",
                                "",
                            )
                        )

                        if (
                            current_partname
                            == image_partname
                        ):

                            image_part = (
                                related_part
                            )

                            break

                # =================================================
                # 3. Image could not be located
                # =================================================

                if image_part is None:

                    element.metadata[
                        "extraction_status"
                    ] = (
                        "failed_image_not_found"
                    )

                    continue

                # =================================================
                # 4. Read original image bytes
                # =================================================

                image_blob = getattr(
                    image_part,
                    "blob",
                    None,
                )

                if not image_blob:

                    element.metadata[
                        "extraction_status"
                    ] = (
                        "failed_empty_image_blob"
                    )

                    continue

                # =================================================
                # 5. Determine original image extension
                # =================================================

                partname = str(
                    getattr(
                        image_part,
                        "partname",
                        "",
                    )
                )

                extension = (
                    Path(
                        partname
                    ).suffix.lower()
                )

                if not extension:

                    extension = ".png"

                # =================================================
                # 6. Save original image
                # =================================================

                image_filename = (
                    f"unit_{unit.index:03d}"
                    f"_element_"
                    f"{element.order:03d}"
                    f"{extension}"
                )

                image_path = (
                    output_path
                    / image_filename
                )

                with open(
                    image_path,
                    "wb",
                ) as image_file:

                    image_file.write(
                        image_blob
                    )

                # =================================================
                # 7. Save extraction result into IR
                # =================================================

                element.metadata[
                    "extracted_image_path"
                ] = str(
                    image_path
                )

                element.metadata[
                    "extracted_image_hash"
                ] = _calculate_extracted_image_hash(
                    image_path
                )

                element.metadata[
                    "extraction_status"
                ] = "done"

                element.metadata[
                    "extracted_extension"
                ] = extension

                extracted_files.append(
                    str(image_path)
                )

    except Exception as e:

        raise RuntimeError(
            f"Failed to extract DOCX images from "
            f"'{document.file_path}': {str(e)}"
        ) from e

    return extracted_files


# =============================================================
# PPTX
# =============================================================


def _extract_pptx_candidate_images(
    document: ProcessedDocument,
    output_dir: str,
    zoom: float = 2.0,
) -> list[str]:
    """
    Extract PPTX visual elements marked with:

        process_with_ai = True

    Strategy:

    1. Normal raster pictures:
       Extract their original embedded bytes directly.

    2. Vector pictures such as WMF / EMF:
       Render the PowerPoint slide and crop the picture area.

    3. Charts and grouped diagrams:
       Render the slide and crop the corresponding shape area.

    Slide rendering is performed through LibreOffice.
    """

    output_path = Path(
        output_dir
    )

    output_path.mkdir(
        parents=True,
        exist_ok=True,
    )

    extracted_files = []

    try:

        presentation = Presentation(
            document.file_path
        )

        candidates = []

        # =====================================================
        # 1. Collect visual candidates
        # =====================================================

        for unit in (
            document.get_ordered_units()
        ):

            for element in (
                unit.get_ordered_elements()
            ):

                if element.type != "image":
                    continue

                if not element.metadata.get(
                    "process_with_ai",
                    False,
                ):
                    continue

                candidates.append(
                    (
                        unit,
                        element,
                    )
                )

        # =====================================================
        # 2. Determine whether slide rendering is required
        # =====================================================

        needs_slide_rendering = False

        for unit, element in candidates:

            visual_type = (
                element.metadata.get(
                    "visual_type",
                    ""
                )
            )

            render_mode = (
                element.metadata.get(
                    "render_mode",
                    ""
                )
            )

            extension = (
                _normalize_extension(
                    element.metadata.get(
                        "image_extension"
                    )
                )
            )

            if (
                render_mode == "slide_crop"
                or visual_type in {
                    "chart",
                    "group",
                }
                or (
                    visual_type == "picture"
                    and extension
                    not in RASTER_IMAGE_EXTENSIONS
                )
            ):

                needs_slide_rendering = True

                break

        # =====================================================
        # 3. Render PPTX to PDF if required
        # =====================================================

        rendered_pdf_path = None
        render_error = None

        if needs_slide_rendering:

            try:

                rendered_pdf_path = (
                    _render_pptx_to_pdf(
                        pptx_path=document.file_path,
                        output_dir=output_path,
                    )
                )

            except Exception as e:

                render_error = str(
                    e
                )

        # =====================================================
        # 4. Open rendered presentation when available
        # =====================================================

        rendered_pdf = None

        if rendered_pdf_path:

            rendered_pdf = fitz.open(
                str(
                    rendered_pdf_path
                )
            )

        try:

            # =================================================
            # 5. Process candidates
            # =================================================

            for unit, element in candidates:

                slide_index = (
                    unit.index - 1
                )

                if (
                    slide_index < 0
                    or slide_index
                    >= len(
                        presentation.slides
                    )
                ):

                    element.metadata[
                        "extraction_status"
                    ] = (
                        "failed_slide_not_found"
                    )

                    continue

                slide = (
                    presentation.slides[
                        slide_index
                    ]
                )

                shape_id = (
                    element.metadata.get(
                        "shape_id"
                    )
                )

                shape = _find_pptx_shape_by_id(
                    slide=slide,
                    shape_id=shape_id,
                )

                if shape is None:

                    element.metadata[
                        "extraction_status"
                    ] = (
                        "failed_shape_not_found"
                    )

                    continue

                visual_type = (
                    element.metadata.get(
                        "visual_type",
                        ""
                    )
                )

                render_mode = (
                    element.metadata.get(
                        "render_mode",
                        ""
                    )
                )

                # =============================================
                # 5A. Embedded picture
                # =============================================

                if (
                    visual_type == "picture"
                    and render_mode == "embedded"
                ):

                    extension = (
                        _get_pptx_picture_extension(
                            shape=shape,
                            element=element,
                        )
                    )

                    # -----------------------------------------
                    # Raster image
                    # -----------------------------------------

                    if (
                        extension
                        in RASTER_IMAGE_EXTENSIONS
                    ):

                        image_blob = (
                            _get_pptx_picture_blob(
                                shape
                            )
                        )

                        if not image_blob:

                            element.metadata[
                                "extraction_status"
                            ] = (
                                "failed_empty_image_blob"
                            )

                            continue

                        image_filename = (
                            f"slide_"
                            f"{unit.index:03d}"
                            f"_element_"
                            f"{element.order:03d}"
                            f"{extension}"
                        )

                        image_path = (
                            output_path
                            / image_filename
                        )

                        with open(
                            image_path,
                            "wb",
                        ) as image_file:

                            image_file.write(
                                image_blob
                            )

                        element.metadata[
                            "extracted_image_path"
                        ] = str(
                            image_path
                        )

                        element.metadata[
                            "extracted_image_hash"
                        ] = _calculate_extracted_image_hash(
                            image_path
                        )

                        element.metadata[
                            "extraction_status"
                        ] = "done"

                        element.metadata[
                            "extracted_extension"
                        ] = extension

                        element.metadata[
                            "extraction_strategy"
                        ] = (
                            "embedded_original"
                        )

                        extracted_files.append(
                            str(image_path)
                        )

                        continue

                    # -----------------------------------------
                    # Vector picture:
                    #
                    # WMF / EMF / SVG / etc.
                    #
                    # Render slide and crop it.
                    # -----------------------------------------

                    success = (
                        _extract_pptx_slide_crop(
                            document=document,
                            unit=unit,
                            element=element,
                            rendered_pdf=rendered_pdf,
                            output_path=output_path,
                            zoom=zoom,
                            original_extension=extension,
                            render_error=render_error,
                        )
                    )

                    if success:

                        extracted_files.append(
                            element.metadata[
                                "extracted_image_path"
                            ]
                        )

                    continue

                # =============================================
                # 5B. Chart / grouped diagram / slide crop
                # =============================================

                if (
                    render_mode == "slide_crop"
                    or visual_type in {
                        "chart",
                        "group",
                    }
                ):

                    success = (
                        _extract_pptx_slide_crop(
                            document=document,
                            unit=unit,
                            element=element,
                            rendered_pdf=rendered_pdf,
                            output_path=output_path,
                            zoom=zoom,
                            original_extension=None,
                            render_error=render_error,
                        )
                    )

                    if success:

                        extracted_files.append(
                            element.metadata[
                                "extracted_image_path"
                            ]
                        )

                    continue

                # =============================================
                # Unknown PPTX visual type
                # =============================================

                element.metadata[
                    "extraction_status"
                ] = (
                    "failed_unsupported_visual_type"
                )

        finally:

            if rendered_pdf is not None:

                rendered_pdf.close()

    except Exception as e:

        raise RuntimeError(
            f"Failed to extract PPTX images from "
            f"'{document.file_path}': {str(e)}"
        ) from e

    return extracted_files


def _extract_pptx_slide_crop(
    document: ProcessedDocument,
    unit,
    element,
    rendered_pdf,
    output_path: Path,
    zoom: float,
    original_extension=None,
    render_error=None,
) -> bool:
    """
    Crop one PPTX visual element from the rendered slide.

    PPTX coordinates are stored in EMU.

    The rendered PDF page uses PDF points.

    Because slide dimensions and element bounding boxes use
    the same PPTX coordinate system, conversion is performed
    using ratios rather than fixed unit conversion.
    """

    if rendered_pdf is None:

        element.metadata[
            "extraction_status"
        ] = (
            "failed_slide_rendering_unavailable"
        )

        if render_error:

            element.metadata[
                "extraction_error"
            ] = render_error

        return False

    slide_index = (
        unit.index - 1
    )

    if (
        slide_index < 0
        or slide_index
        >= rendered_pdf.page_count
    ):

        element.metadata[
            "extraction_status"
        ] = (
            "failed_rendered_slide_not_found"
        )

        return False

    if element.bbox is None:

        element.metadata[
            "extraction_status"
        ] = (
            "failed_missing_bbox"
        )

        return False

    slide_width = float(
        unit.metadata.get(
            "width",
            0,
        )
    )

    slide_height = float(
        unit.metadata.get(
            "height",
            0,
        )
    )

    if (
        slide_width <= 0
        or slide_height <= 0
    ):

        element.metadata[
            "extraction_status"
        ] = (
            "failed_invalid_slide_dimensions"
        )

        return False

    page = rendered_pdf[
        slide_index
    ]

    bbox = element.bbox

    # =========================================================
    # Convert PPTX bbox to ratios
    # =========================================================

    x_ratio = (
        bbox.x
        / slide_width
    )

    y_ratio = (
        bbox.y
        / slide_height
    )

    width_ratio = (
        bbox.width
        / slide_width
    )

    height_ratio = (
        bbox.height
        / slide_height
    )

    # =========================================================
    # Convert ratios into rendered PDF coordinates
    # =========================================================

    x0 = (
        page.rect.x0
        + (
            x_ratio
            * page.rect.width
        )
    )

    y0 = (
        page.rect.y0
        + (
            y_ratio
            * page.rect.height
        )
    )

    x1 = (
        x0
        + (
            width_ratio
            * page.rect.width
        )
    )

    y1 = (
        y0
        + (
            height_ratio
            * page.rect.height
        )
    )

    rect = fitz.Rect(
        x0,
        y0,
        x1,
        y1,
    )

    rect = (
        rect
        & page.rect
    )

    if rect.is_empty:

        element.metadata[
            "extraction_status"
        ] = (
            "failed_empty_render_crop"
        )

        return False

    matrix = fitz.Matrix(
        zoom,
        zoom,
    )

    pixmap = page.get_pixmap(
        matrix=matrix,
        clip=rect,
        alpha=False,
    )

    image_filename = (
        f"slide_{unit.index:03d}"
        f"_element_"
        f"{element.order:03d}.png"
    )

    image_path = (
        output_path
        / image_filename
    )

    pixmap.save(
        str(image_path)
    )

    # =========================================================
    # Store extraction information
    # =========================================================

    element.metadata[
        "extracted_image_path"
    ] = str(
        image_path
    )

    element.metadata[
        "extracted_image_hash"
    ] = _calculate_extracted_image_hash(
        image_path
    )

    element.metadata[
        "extraction_status"
    ] = "done"

    element.metadata[
        "extracted_extension"
    ] = ".png"

    element.metadata[
        "extraction_strategy"
    ] = "slide_crop"

    element.metadata[
        "render_zoom"
    ] = zoom

    if original_extension:

        element.metadata[
            "original_extension"
        ] = original_extension

    element.metadata.pop(
        "extraction_error",
        None,
    )

    return True



def _calculate_extracted_image_hash(
    image_path: Path,
) -> str:
    """
    Calculate SHA256 for the final extracted image file.

    This hash represents the exact image file that will
    later be sent to OCR / Vision.

    The hash is calculated after extraction/rendering so
    the same mechanism works consistently for:
    - PDF crops
    - DOCX embedded images
    - PPTX embedded raster images
    - PPTX slide_crop images
    """

    sha256 = hashlib.sha256()

    with open(
        image_path,
        "rb",
    ) as image_file:

        while True:

            chunk = image_file.read(
                1024 * 1024
            )

            if not chunk:
                break

            sha256.update(
                chunk
            )

    return sha256.hexdigest()

def _find_pptx_shape_by_id(
    slide,
    shape_id,
):
    """
    Find a top-level PowerPoint shape using its shape_id.
    """

    if shape_id is None:
        return None

    for shape in slide.shapes:

        if (
            shape.shape_id
            == shape_id
        ):

            return shape

    return None


def _get_pptx_picture_blob(
    shape,
):
    """
    Safely return the embedded picture bytes from
    a normal Picture or PlaceholderPicture.
    """

    try:

        image = shape.image

        return image.blob

    except Exception:

        return None


def _get_pptx_picture_extension(
    shape,
    element,
) -> str:
    """
    Determine a normalized PowerPoint picture extension.

    Example:

        png  -> .png
        jpg  -> .jpg
        wmf  -> .wmf
    """

    extension = ""

    try:

        image = shape.image

        extension = getattr(
            image,
            "ext",
            "",
        )

    except Exception:

        extension = ""

    if not extension:

        extension = (
            element.metadata.get(
                "image_extension",
                "",
            )
        )

    return _normalize_extension(
        extension
    )


def _normalize_extension(
    extension,
) -> str:
    """
    Normalize image extensions into lowercase
    strings beginning with a dot.
    """

    if not extension:
        return ""

    extension = str(
        extension
    ).strip().lower()

    if not extension:
        return ""

    if not extension.startswith(
        "."
    ):

        extension = (
            "."
            + extension
        )

    return extension


def _render_pptx_to_pdf(
    pptx_path: str,
    output_dir: Path,
) -> Path:
    """
    Render the PPTX presentation to PDF using LibreOffice.

    This allows:
    - WMF / EMF pictures
    - charts
    - grouped diagrams

    to be converted into normal raster crops later.

    The function searches:
    1. PATH
    2. Common Windows LibreOffice locations
    """

    libreoffice = (
        _find_libreoffice_executable()
    )

    if libreoffice is None:

        raise RuntimeError(
            "LibreOffice was not found. "
            "PPTX vector pictures, charts, and grouped "
            "diagrams require slide rendering before "
            "OCR / Vision."
        )

    render_dir = (
        output_dir
        / "_pptx_render"
    )

    render_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    command = [
        str(libreoffice),
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(render_dir),
        str(
            Path(
                pptx_path
            ).resolve()
        ),
    ]

    result = subprocess.run(
        command,
        capture_output=True,
        text=True,
        timeout=120,
    )

    expected_pdf = (
        render_dir
        / (
            Path(
                pptx_path
            ).stem
            + ".pdf"
        )
    )

    if (
        result.returncode != 0
        or not expected_pdf.exists()
    ):

        stdout = (
            result.stdout
            or ""
        ).strip()

        stderr = (
            result.stderr
            or ""
        ).strip()

        raise RuntimeError(
            "LibreOffice failed to render PPTX. "
            f"stdout='{stdout}' "
            f"stderr='{stderr}'"
        )

    return expected_pdf


def _find_libreoffice_executable():
    """
    Find LibreOffice / soffice executable.

    Supports:
    - Linux/macOS PATH
    - Windows PATH
    - common Windows installation locations
    """

    # =========================================================
    # PATH
    # =========================================================

    for command in (
        "soffice",
        "libreoffice",
    ):

        executable = shutil.which(
            command
        )

        if executable:

            return Path(
                executable
            )

    # =========================================================
    # Common Windows locations
    # =========================================================

    windows_candidates = [
        Path(
            r"C:\Program Files\LibreOffice"
            r"\program\soffice.exe"
        ),
        Path(
            r"C:\Program Files (x86)\LibreOffice"
            r"\program\soffice.exe"
        ),
    ]

    for candidate in windows_candidates:

        if candidate.exists():

            return candidate

    return None