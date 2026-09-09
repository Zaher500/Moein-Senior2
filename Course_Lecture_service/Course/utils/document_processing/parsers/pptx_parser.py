import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..schemas import (
    BoundingBox,
    DocumentElement,
    DocumentUnit,
    ProcessedDocument,
)
from ..image_filter import mark_image_candidates


def parse_pptx(
    file_path: str,
) -> ProcessedDocument:
    """
    Parse a PPTX file into the unified document IR.

    Each PowerPoint slide becomes one DocumentUnit.

    Supported:
    - Native text
    - Tables
    - Embedded pictures
    - Pictures inside placeholders
    - Charts
    - Grouped visual elements

    Charts, grouped visuals, and malformed/unextractable
    picture objects are rendered later from the slide
    using the slide_crop extraction path.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"PPTX file not found: {file_path}"
        )

    presentation = Presentation(
        str(path)
    )

    document = ProcessedDocument(
        file_path=str(path),
        file_type="pptx",
    )

    # =========================================================
    # Document metadata
    # =========================================================

    properties = (
        presentation.core_properties
    )

    document.metadata.update({
        "slide_count": len(
            presentation.slides
        ),
        "title": properties.title or "",
        "author": properties.author or "",
        "subject": properties.subject or "",
    })

    slide_width = float(
        presentation.slide_width
    )

    slide_height = float(
        presentation.slide_height
    )

    # =========================================================
    # Parse slides
    # =========================================================

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):

        unit = DocumentUnit(
            index=slide_number,
            unit_type="slide",
            metadata={
                "width": slide_width,
                "height": slide_height,
            },
        )

        slide_elements = []

        # =====================================================
        # Parse shapes
        # =====================================================

        for shape in slide.shapes:

            elements = _parse_shape(
                shape=shape,
                slide_number=slide_number,
            )

            slide_elements.extend(
                elements
            )

        # =====================================================
        # Approximate visual reading order
        # =====================================================

        slide_elements.sort(
            key=_element_sort_key
        )

        # =====================================================
        # Assign final order
        # =====================================================

        for order, element in enumerate(
            slide_elements,
            start=1,
        ):

            element.order = order

            unit.add_element(
                element
            )

        document.add_unit(
            unit
        )

    # =========================================================
    # Mark PPTX image candidates only once, after the whole
    # presentation has been parsed.
    # =========================================================

    mark_image_candidates(
        document
    )

    return document


def _parse_shape(
    shape,
    slide_number: int,
) -> list[DocumentElement]:
    """
    Convert one PowerPoint shape into one or more
    DocumentElement objects.
    """

    elements = []

    # =========================================================
    # 1. Table
    # =========================================================

    if getattr(
        shape,
        "has_table",
        False,
    ):

        table_element = (
            _build_table_element(
                shape=shape,
                slide_number=slide_number,
            )
        )

        if table_element:
            elements.append(
                table_element
            )

        return elements

    # =========================================================
    # 2. Embedded picture / picture placeholder
    # =========================================================

    if _is_picture_shape(
        shape
    ):

        picture_element = (
            _build_picture_element(
                shape=shape,
                slide_number=slide_number,
            )
        )

        if picture_element:
            elements.append(
                picture_element
            )

        return elements

    # =========================================================
    # 3. Chart
    # =========================================================

    if getattr(
        shape,
        "has_chart",
        False,
    ):

        chart_element = (
            _build_chart_element(
                shape=shape,
                slide_number=slide_number,
            )
        )

        elements.append(
            chart_element
        )

        return elements

    # =========================================================
    # 4. Grouped shapes / diagrams
    # =========================================================

    if (
        shape.shape_type
        == MSO_SHAPE_TYPE.GROUP
    ):

        # ---------------------------------------------
        # Preserve native text found inside the group
        # ---------------------------------------------

        group_text = (
            _extract_group_text(
                shape
            )
        )

        if group_text:

            elements.append(
                DocumentElement(
                    order=0,
                    type="text",
                    source="native",
                    text=group_text,
                    bbox=_shape_bbox(
                        shape
                    ),
                    metadata={
                        "slide_number":
                            slide_number,
                        "shape_id":
                            shape.shape_id,
                        "shape_name":
                            shape.name,
                        "content_type":
                            "group_text",
                    },
                )
            )

        # ---------------------------------------------
        # Preserve the full group as a visual candidate
        # ---------------------------------------------

        elements.append(
            DocumentElement(
                order=0,
                type="image",
                source="structured",
                bbox=_shape_bbox(
                    shape
                ),
                metadata={
                    "slide_number":
                        slide_number,
                    "shape_id":
                        shape.shape_id,
                    "shape_name":
                        shape.name,
                    "visual_type":
                        "group",
                    "render_mode":
                        "slide_crop",
                },
            )
        )

        return elements

    # =========================================================
    # 5. Normal native text
    # =========================================================

    if getattr(
        shape,
        "has_text_frame",
        False,
    ):

        text = _clean_text(
            shape.text
        )

        if text:

            elements.append(
                DocumentElement(
                    order=0,
                    type="text",
                    source="native",
                    text=text,
                    bbox=_shape_bbox(
                        shape
                    ),
                    metadata={
                        "slide_number":
                            slide_number,
                        "shape_id":
                            shape.shape_id,
                        "shape_name":
                            shape.name,
                        "content_type":
                            "text",
                    },
                )
            )

    return elements


def _is_picture_shape(
    shape,
) -> bool:
    """
    Detect PowerPoint pictures.

    Supports:
    - Normal picture shapes
    - Pictures stored inside placeholders
    - Malformed picture placeholders where python-pptx
      recognizes the picture object but cannot expose
      shape.image because <p:blipFill> is missing.

    A malformed picture is still treated as a visual so
    _build_picture_element() can safely route it to
    slide_crop instead of crashing the whole lecture.
    """

    # =========================================================
    # Normal picture
    # =========================================================

    if (
        shape.shape_type
        == MSO_SHAPE_TYPE.PICTURE
    ):
        return True

    # =========================================================
    # Picture inside placeholder
    # =========================================================

    if getattr(
        shape,
        "is_placeholder",
        False,
    ):

        try:

            image = shape.image

            if image is not None:
                return True

        except AttributeError:
            # Normal non-picture placeholder.
            return False

        except ValueError as exc:
            # Some malformed picture placeholders are still
            # visual picture objects, but python-pptx cannot
            # expose shape.image because <p:blipFill> is absent.
            #
            # Keep them as pictures so the next step can use
            # slide_crop instead of losing the visual.
            error_message = str(exc)

            if "blipFill" in error_message:
                return True

            return False

    return False


def _build_picture_element(
    shape,
    slide_number: int,
) -> DocumentElement:
    """
    Build an image element from a PowerPoint picture.

    Normal case:
        shape.image is readable
        -> render_mode = embedded

    Fallback case:
        shape.image cannot be read, for example:
        "required <p:blipFill> child element not present"
        -> render_mode = slide_crop

    This prevents one malformed picture from crashing
    the entire PPTX lecture while still preserving the
    visual for OCR / Vision.
    """

    # =========================================================
    # Try normal embedded-image extraction metadata first
    # =========================================================

    try:

        image = shape.image

        blob = image.blob

        content_hash = hashlib.sha256(
            blob
        ).hexdigest()

        extension = (
            getattr(
                image,
                "ext",
                None,
            )
            or ""
        )

        content_type = (
            getattr(
                image,
                "content_type",
                None,
            )
            or ""
        )

        return DocumentElement(
            order=0,
            type="image",
            source="structured",
            bbox=_shape_bbox(
                shape
            ),
            metadata={
                "slide_number":
                    slide_number,
                "shape_id":
                    shape.shape_id,
                "shape_name":
                    shape.name,
                "content_hash":
                    content_hash,
                "image_extension":
                    extension,
                "content_type":
                    content_type,
                "visual_type":
                    "picture",
                "render_mode":
                    "embedded",
                "is_placeholder":
                    bool(
                        getattr(
                            shape,
                            "is_placeholder",
                            False,
                        )
                    ),
            },
        )

    except Exception as exc:

        # =====================================================
        # Fallback:
        #
        # python-pptx could not expose the embedded image.
        # Do NOT fail the whole lecture and do NOT discard
        # the visual. Render its slide region later instead.
        # =====================================================

        error_message = str(exc)

        print(
            "[PPTX PARSER]",
            "PICTURE_FALLBACK_TO_SLIDE_CROP",
            f"slide={slide_number}",
            f"shape_id={getattr(shape, 'shape_id', '')}",
            f"shape_name={getattr(shape, 'name', '')}",
            f"error={error_message}",
        )

        return DocumentElement(
            order=0,
            type="image",
            source="structured",
            bbox=_shape_bbox(
                shape
            ),
            metadata={
                "slide_number":
                    slide_number,
                "shape_id":
                    shape.shape_id,
                "shape_name":
                    shape.name,
                "visual_type":
                    "picture",
                "render_mode":
                    "slide_crop",
                "is_placeholder":
                    bool(
                        getattr(
                            shape,
                            "is_placeholder",
                            False,
                        )
                    ),
                "embedded_image_unavailable":
                    True,
                "embedded_image_error":
                    error_message,
                "fallback_reason":
                    "embedded_picture_unavailable",
            },
        )


def _build_chart_element(
    shape,
    slide_number: int,
) -> DocumentElement:
    """
    Preserve a PowerPoint chart as a visual element.

    Charts cannot be extracted through shape.image.

    They will later be rendered from the slide
    and passed through OCR / Vision.
    """

    chart_type = ""

    try:

        chart_type = str(
            shape.chart.chart_type
        )

    except Exception:

        chart_type = ""

    return DocumentElement(
        order=0,
        type="image",
        source="structured",
        bbox=_shape_bbox(
            shape
        ),
        metadata={
            "slide_number":
                slide_number,
            "shape_id":
                shape.shape_id,
            "shape_name":
                shape.name,
            "visual_type":
                "chart",
            "render_mode":
                "slide_crop",
            "chart_type":
                chart_type,
        },
    )


def _build_table_element(
    shape,
    slide_number: int,
):
    """
    Convert a native PowerPoint table into readable text.

    Example:

        Epoch | Loss | Accuracy
        1 | 0.32 | 89%
        2 | 0.08 | 97%
    """

    table = shape.table

    rows = []

    for row in table.rows:

        row_values = []

        for cell in row.cells:

            value = _clean_text(
                cell.text
            )

            value = value.replace(
                "\n",
                " ",
            )

            row_values.append(
                value
            )

        if any(
            value
            for value in row_values
        ):

            rows.append(
                " | ".join(
                    row_values
                )
            )

    table_text = "\n".join(
        rows
    ).strip()

    if not table_text:
        return None

    return DocumentElement(
        order=0,
        type="text",
        source="native",
        text=table_text,
        bbox=_shape_bbox(
            shape
        ),
        metadata={
            "slide_number":
                slide_number,
            "shape_id":
                shape.shape_id,
            "shape_name":
                shape.name,
            "content_type":
                "table",
            "row_count":
                len(table.rows),
            "column_count":
                len(table.columns),
        },
    )


def _extract_group_text(
    group_shape,
) -> str:
    """
    Extract readable native text from shapes inside
    a PowerPoint group.

    We preserve approximate local visual order.
    """

    group_items = []

    for child in group_shape.shapes:

        # Nested group
        if (
            child.shape_type
            == MSO_SHAPE_TYPE.GROUP
        ):

            nested_text = (
                _extract_group_text(
                    child
                )
            )

            if nested_text:

                group_items.append(
                    (
                        _safe_position(
                            child,
                            "top",
                        ),
                        _safe_position(
                            child,
                            "left",
                        ),
                        nested_text,
                    )
                )

            continue

        # Table inside group
        if getattr(
            child,
            "has_table",
            False,
        ):

            rows = []

            for row in child.table.rows:

                values = [
                    _clean_text(
                        cell.text
                    ).replace(
                        "\n",
                        " ",
                    )
                    for cell in row.cells
                ]

                if any(values):

                    rows.append(
                        " | ".join(
                            values
                        )
                    )

            table_text = "\n".join(
                rows
            ).strip()

            if table_text:

                group_items.append(
                    (
                        _safe_position(
                            child,
                            "top",
                        ),
                        _safe_position(
                            child,
                            "left",
                        ),
                        table_text,
                    )
                )

            continue

        # Normal text child
        if getattr(
            child,
            "has_text_frame",
            False,
        ):

            text = _clean_text(
                child.text
            )

            if text:

                group_items.append(
                    (
                        _safe_position(
                            child,
                            "top",
                        ),
                        _safe_position(
                            child,
                            "left",
                        ),
                        text,
                    )
                )

    group_items.sort(
        key=lambda item: (
            item[0],
            item[1],
        )
    )

    return "\n".join(
        item[2]
        for item in group_items
    ).strip()


def _shape_bbox(
    shape,
) -> BoundingBox:
    """
    Build BoundingBox using PowerPoint EMU coordinates.

    Slide dimensions are also stored using the same unit,
    so ratios can later be calculated safely.
    """

    return BoundingBox(
        x=float(
            getattr(
                shape,
                "left",
                0,
            )
            or 0
        ),
        y=float(
            getattr(
                shape,
                "top",
                0,
            )
            or 0
        ),
        width=float(
            getattr(
                shape,
                "width",
                0,
            )
            or 0
        ),
        height=float(
            getattr(
                shape,
                "height",
                0,
            )
            or 0
        ),
    )


def _element_sort_key(
    element: DocumentElement,
):
    """
    Approximate human reading order:
    top -> left.

    Native text is placed before a visual element
    if both occupy the same position.
    """

    if element.bbox is None:

        return (
            float("inf"),
            float("inf"),
            0,
        )

    type_priority = (
        0
        if element.type == "text"
        else 1
    )

    return (
        element.bbox.y,
        element.bbox.x,
        type_priority,
    )


def _safe_position(
    shape,
    attribute: str,
) -> float:

    try:

        return float(
            getattr(
                shape,
                attribute,
                0,
            )
            or 0
        )

    except Exception:

        return 0.0


def _clean_text(
    text,
) -> str:

    if not text:
        return ""

    return (
        str(text)
        .replace(
            "\x00",
            "",
        )
        .replace(
            "\u0000",
            "",
        )
        .strip()
    )
